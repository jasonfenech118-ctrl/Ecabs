from datetime import timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from .models import Claim, OtherCharge, Reminder, Vehicle, compliance_state


class ClaimModelTests(TestCase):
    def test_reference_sequence(self):
        a = Claim.objects.create()
        b = Claim.objects.create()
        self.assertTrue(a.reference.startswith("CLM-"))
        self.assertEqual(
            int(b.reference.rsplit("-", 1)[1]),
            int(a.reference.rsplit("-", 1)[1]) + 1,
        )

    def test_submit_promotes_draft(self):
        claim = Claim.objects.create()
        self.assertTrue(claim.is_draft)
        claim.submit()
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.OPEN)
        self.assertIsNotNone(claim.submitted_at)

    def test_history_recorded(self):
        claim = Claim.objects.create()
        claim.vehicle_registration = "ECB-001"
        claim.save()
        self.assertEqual(claim.history.count(), 2)


class ViewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("staff", password="pw")
        self.client.force_login(self.user)

    def test_dashboard(self):
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)

    def test_claim_new_creates_draft_and_redirects(self):
        response = self.client.post(reverse("claim_new"))
        claim = Claim.objects.get()
        self.assertRedirects(response, reverse("claim_edit", args=[claim.pk]))
        self.assertTrue(claim.is_draft)

    def test_autosave_updates_draft(self):
        claim = Claim.objects.create(created_by=self.user)
        response = self.client.post(
            reverse("claim_autosave", args=[claim.pk]),
            {"vehicle_registration": "ECB-777", "fault": "unknown"},
        )
        self.assertEqual(response.status_code, 200)
        claim.refresh_from_db()
        self.assertEqual(claim.vehicle_registration, "ECB-777")
        self.assertTrue(claim.is_draft)

    def test_search_returns_partial_for_htmx(self):
        Claim.objects.create(vehicle_registration="ECB-999", created_by=self.user)
        response = self.client.get(
            reverse("claim_list"), {"q": "ECB-999"}, headers={"HX-Request": "true"}
        )
        self.assertContains(response, "ECB-999")
        self.assertNotContains(response, "<html")

    def test_tab_partial(self):
        claim = Claim.objects.create(created_by=self.user)
        response = self.client.get(
            reverse("claim_tab", args=[claim.pk, "billing"]),
            headers={"HX-Request": "true"},
        )
        self.assertEqual(response.status_code, 200)

    def test_reminder_toggle(self):
        claim = Claim.objects.create(created_by=self.user)
        reminder = Reminder.objects.create(
            claim=claim, title="Chase insurer", due_at="2026-01-01T10:00:00Z"
        )
        self.client.post(reverse("reminder_toggle", args=[reminder.pk]))
        reminder.refresh_from_db()
        self.assertTrue(reminder.is_done)

    def test_recovery_totals(self):
        from decimal import Decimal

        claim = Claim.objects.create(
            labour_amount=Decimal("243.60"),
            spray_material_amount=Decimal("429.53"),
            parts_amount=Decimal("800.21"),
            loe_days=4,
            loe_daily_rate=Decimal("46.405"),
            amount_paid=Decimal("1023.34"),
            offset_amount=Decimal("240.00"),
        )
        self.assertEqual(claim.loss_of_earnings, Decimal("185.62"))
        self.assertEqual(claim.total_claim_amount, Decimal("1658.96"))
        self.assertEqual(claim.outstanding_amount, Decimal("395.62"))

    def test_search_by_tp_claim_number(self):
        Claim.objects.create(tp_claim_number="C34-277148", created_by=self.user)
        response = self.client.get(
            reverse("claim_list"), {"q": "C34-277148"}, headers={"HX-Request": "true"}
        )
        self.assertContains(response, "CLM-")

    def test_urgent_claim_on_dashboard_chasers(self):
        claim = Claim.objects.create(
            status=Claim.Status.OPEN,
            vehicle_registration="ECB-800",
            urgent=True,
            next_action="O/S payment from MSI",
        )
        response = self.client.get(reverse("dashboard"))
        self.assertContains(response, claim.reference)
        self.assertContains(response, "O/S payment from MSI")

    def test_vehicle_add_retire_delete(self):
        response = self.client.post(
            reverse("vehicle_add"),
            {"registration": "ECB-500", "make_model": "Toyota Yaris"},
        )
        vehicle = Vehicle.objects.get(registration="ECB-500")
        self.assertRedirects(response, reverse("vehicle_edit", args=[vehicle.pk]))
        self.assertTrue(vehicle.is_active)

        self.client.post(reverse("vehicle_toggle_status", args=[vehicle.pk]))
        vehicle.refresh_from_db()
        self.assertFalse(vehicle.is_active)
        self.assertIsNotNone(vehicle.retired_on)

        self.client.post(reverse("vehicle_delete", args=[vehicle.pk]))
        self.assertFalse(Vehicle.objects.filter(pk=vehicle.pk).exists())

    def test_compliance_states_and_dashboard_reminder(self):
        today = timezone.localdate()
        self.assertEqual(compliance_state(today - timedelta(days=1)), "overdue")
        self.assertEqual(compliance_state(today + timedelta(days=10)), "soon")
        self.assertEqual(compliance_state(today + timedelta(days=90)), "ok")
        self.assertEqual(compliance_state(None), "")

        from .models import VehicleCost
        v600 = Vehicle.objects.create(registration="ECB-600")
        VehicleCost.objects.create(vehicle=v600, year=today.year,
                                   pay_date=today + timedelta(days=5))
        response = self.client.get(reverse("dashboard"))
        self.assertContains(response, "ECB-600")
        self.assertContains(response, "due soon")

    def test_retired_vehicle_not_in_compliance_list(self):
        from .models import VehicleCost
        today = timezone.localdate()
        v700 = Vehicle.objects.create(
            registration="ECB-700", status=Vehicle.Status.RETIRED)
        VehicleCost.objects.create(vehicle=v700, year=today.year,
                                   pay_date=today - timedelta(days=5))
        response = self.client.get(reverse("dashboard"))
        self.assertNotContains(response, "ECB-700")

    def test_home_page_tiles(self):
        response = self.client.get(reverse("home"))
        self.assertContains(response, "New claim")
        self.assertContains(response, "Reminders")

    def test_vehicle_yearly_costs_and_owner_grouping(self):
        from decimal import Decimal

        from .models import Vehicle, VehicleCost

        v = Vehicle.objects.create(registration="OWN-1", owner="Vai Drive Co Ltd")
        # add two years of costs
        self.client.post(reverse("vehicle_cost_save", args=[v.pk]),
                         {"year": "2025", "insurance_amount": "400",
                          "licence_amount": "100", "additional_costs": "50"})
        self.client.post(reverse("vehicle_cost_save", args=[v.pk]),
                         {"year": "2026", "insurance_amount": "440",
                          "licence_amount": "110"})
        self.assertEqual(v.costs.count(), 2)
        # current (latest year) costs feed the vehicle
        self.assertEqual(v.current_cost.year, 2026)
        self.assertEqual(v.total_costs, Decimal("550.00"))  # 440 + 110

        # fleet list groups by owner and shows the current total
        r = self.client.get(reverse("vehicle_list"))
        self.assertContains(r, "Vai Drive Co Ltd")
        self.assertContains(r, "OWN-1")

        # deleting a year works
        self.client.post(reverse("vehicle_cost_delete", args=[v.pk, 2025]))
        self.assertEqual(v.costs.count(), 1)

    def test_multiple_additional_costs_with_note_and_date(self):
        from decimal import Decimal

        from .models import Vehicle, VehicleCost

        v = Vehicle.objects.create(registration="ADD-2", owner="Fast Drop")
        self.client.post(reverse("vehicle_cost_save", args=[v.pk]),
                         {"year": "2026", "insurance_amount": "440", "licence_amount": "110"})
        cost = VehicleCost.objects.get(vehicle=v, year=2026)
        # two additional items, each with a note and a date
        self.client.post(reverse("vehicle_additional_add", args=[cost.pk]),
                         {"amount": "50", "note": "new tyres", "date_added": "2026-08-12"})
        self.client.post(reverse("vehicle_additional_add", args=[cost.pk]),
                         {"amount": "30", "note": "wiper", "date_added": "2026-08-15"})
        cost.refresh_from_db()
        self.assertEqual(cost.additionals.count(), 2)
        self.assertEqual(cost.additional_total, Decimal("80.00"))
        self.assertEqual(cost.total, Decimal("630.00"))  # 440 + 110 + 80
        item = cost.additionals.first()
        self.assertEqual(item.note, "new tyres")
        # delete one
        self.client.post(reverse("vehicle_additional_delete", args=[cost.pk, item.pk]))
        cost.refresh_from_db()
        self.assertEqual(cost.additionals.count(), 1)

    def test_standalone_reminder(self):
        response = self.client.post(
            reverse("reminder_create"),
            {"title": "Call surveyor", "due_at": "2026-08-01T10:00"},
        )
        self.assertRedirects(response, reverse("reminder_list"))
        reminder = Reminder.objects.get(title="Call surveyor")
        self.assertIsNone(reminder.claim)
        self.assertFalse(reminder.is_auto)

    def test_auto_reminder_for_insurance_and_chase(self):
        from .services.auto_reminders import sync_auto_reminders

        from .models import VehicleCost
        today = timezone.localdate()
        vehicle = Vehicle.objects.create(registration="ECB-810")
        vcost = VehicleCost.objects.create(vehicle=vehicle, year=today.year,
                                           pay_date=today + timedelta(days=10))
        claim = Claim.objects.create(
            status=Claim.Status.OPEN,
            chase_on=today + timedelta(days=3),
            next_action="O/S payment from MSI",
        )
        sync_auto_reminders()
        sync_auto_reminders()  # idempotent — must not duplicate
        autos = Reminder.objects.filter(is_auto=True)
        self.assertEqual(autos.count(), 2)
        self.assertTrue(autos.filter(title__contains="ECB-810", claim__isnull=True).exists())
        self.assertTrue(autos.filter(claim=claim, title__contains="O/S payment").exists())

        # Date moves on -> stale auto reminder replaced, not duplicated.
        vcost.pay_date = today + timedelta(days=40)
        vcost.save()
        sync_auto_reminders()
        self.assertEqual(Reminder.objects.filter(is_auto=True).count(), 2)

    def test_bills_report_month_and_csv(self):
        from datetime import date
        from decimal import Decimal

        Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-BILL",
            third_party_insurer="Mapfre Middlesea",
            bills_sent_on=date(2026, 7, 5), invoice_number="100050",
            labour_amount=Decimal("100.00"), parts_amount=Decimal("50.00"),
            created_by=self.user,
        )
        Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-OTHER",
            bills_sent_on=date(2026, 6, 20), created_by=self.user,
        )
        r = self.client.get(reverse("bills_report"), {"month": "2026-07"})
        self.assertContains(r, "ECB-BILL")
        self.assertNotContains(r, "ECB-OTHER")
        self.assertContains(r, "July 2026")
        self.assertContains(r, "Mapfre Middlesea")  # grouped by insurer
        self.assertContains(r, "150.00")  # total for the month

        csv = self.client.get(reverse("bills_report"), {"month": "2026-07", "format": "csv"})
        self.assertEqual(csv["Content-Type"], "text/csv")
        self.assertIn("ECB-BILL", csv.content.decode())
        self.assertIn("100050", csv.content.decode())

    def test_claim_group_tables(self):
        from datetime import timedelta

        today = timezone.localdate()
        open_c = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-OPN", created_by=self.user
        )
        closed_c = Claim.objects.create(
            status=Claim.Status.CLOSED, vehicle_registration="ECB-CLS", created_by=self.user
        )
        overdue_c = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-OVD",
            chase_on=today - timedelta(days=3), created_by=self.user,
        )

        r = self.client.get(reverse("claim_group", args=["open"]))
        self.assertContains(r, "ECB-OPN")
        self.assertContains(r, "ECB-OVD")  # overdue is also open
        self.assertNotContains(r, "ECB-CLS")

        r = self.client.get(reverse("claim_group", args=["closed"]))
        self.assertContains(r, "ECB-CLS")
        self.assertNotContains(r, "ECB-OPN")

        r = self.client.get(reverse("claim_group", args=["overdue"]))
        self.assertContains(r, "ECB-OVD")
        self.assertNotContains(r, "ECB-OPN")

        self.assertEqual(
            self.client.get(reverse("claim_group", args=["bogus"])).status_code, 400
        )

    def test_open_group_includes_custom_status(self):
        """A claim moved to a custom status is still 'open' (not draft/finished)."""
        from claims.models import ClaimStatus

        cs = ClaimStatus.objects.create(name="O/S To open a claim")
        custom_c = Claim.objects.create(
            status=cs.slug, vehicle_registration="ECB-CST", created_by=self.user
        )
        r = self.client.get(reverse("claim_group", args=["open"]))
        self.assertContains(r, "ECB-CST")

    def test_invoice_company_selection(self):
        from decimal import Decimal

        from .models import Company

        co = Company.objects.create(
            name="Vai Drive Co Ltd.", address="Paceville, Malta",
            bank_name="Banif Bank", iban="MT40BNIF...",
            logo_static="img/companies/vai.png",
        )
        claim = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-INV",
            third_party_registration="ABC123", third_party_insurer="Mapfre",
            tp_claim_number="C34-999", labour_amount=Decimal("100.00"),
            parts_amount=Decimal("250.00"), loe_days=3,
            loe_daily_rate=Decimal("35.00"), created_by=self.user,
        )
        # No company chosen yet -> prompt, no summary.
        r = self.client.get(reverse("claim_documents", args=[claim.pk]), {"doc": "statement"})
        self.assertContains(r, "Select a company")
        self.assertNotContains(r, "Total NET")
        # Company chosen -> statement summary with letterhead + total.
        r = self.client.get(
            reverse("claim_documents", args=[claim.pk]),
            {"doc": "statement", "company": co.pk},
        )
        self.assertContains(r, "Vai Drive Co Ltd.")
        self.assertContains(r, "455.00")  # 100 + 250 + 3*35
        self.assertContains(r, "Total NET")

    def test_maintenance_rates(self):
        from decimal import Decimal

        from .models import DailyRate

        # Add via the maintenance page
        self.client.post(reverse("maintenance"), {"name": "Van", "amount": "45.00"})
        rate = DailyRate.objects.get(name="Van")
        self.assertEqual(rate.amount, Decimal("45.00"))
        # Edit
        self.client.post(
            reverse("rate_edit", args=[rate.pk]),
            {"name": "Van (large)", "amount": "50.00", "is_active": "on"},
        )
        rate.refresh_from_db()
        self.assertEqual(rate.name, "Van (large)")
        self.assertEqual(rate.amount, Decimal("50.00"))
        # The claim form offers the maintained rate
        claim = Claim.objects.create(created_by=self.user)
        r = self.client.get(reverse("claim_edit", args=[claim.pk]))
        self.assertContains(r, "Or pick a saved rate")
        self.assertContains(r, "Van (large)")
        # Delete
        self.client.post(reverse("rate_delete", args=[rate.pk]))
        self.assertFalse(DailyRate.objects.filter(pk=rate.pk).exists())

    def test_other_charges_lines(self):
        from decimal import Decimal

        claim = Claim.objects.create(created_by=self.user)
        # Autosave posts two "other" lines
        self.client.post(
            reverse("claim_autosave", args=[claim.pk]),
            {
                "fault": "unknown",
                "other_desc": ["Wheel alignment", "Windscreen repair", ""],
                "other_amt": ["45.00", "120.50", ""],
            },
        )
        claim.refresh_from_db()
        self.assertEqual(claim.other_charges.count(), 2)  # blank row ignored
        self.assertEqual(claim.total_claim_amount, Decimal("165.50"))
        # They appear as invoice lines with their descriptions
        items = [line[0] for line in claim.invoice_lines()]
        self.assertIn("Wheel alignment", items)
        self.assertIn("Windscreen repair", items)
        # Re-posting replaces, not appends
        self.client.post(
            reverse("claim_autosave", args=[claim.pk]),
            {"fault": "unknown", "other_desc": ["Towing"], "other_amt": ["80.00"]},
        )
        claim.refresh_from_db()
        self.assertEqual(claim.other_charges.count(), 1)
        self.assertEqual(claim.total_claim_amount, Decimal("80.00"))

    def test_repairs_invoice_pdf(self):
        from decimal import Decimal

        from .models import Company

        co = Company.objects.create(
            name="eCabs Ltd", address="St Julians, Malta",
            vat_no="MT21583611", exo_number="4270",
            logo_static="img/companies/ecabs.png",
        )
        claim = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="BLY438",
            tp_claim_number="M24109061", labour_amount=Decimal("200.00"),
            spray_material_amount=Decimal("168.02"), created_by=self.user,
        )
        # net = 368.02, vat @18% = 66.24, total = 434.26
        self.assertEqual(claim.repairs_total, Decimal("368.02"))
        r = self.client.get(
            reverse("claim_documents", args=[claim.pk]),
            {"doc": "repairs", "company": co.pk},
        )
        self.assertContains(r, "VAT @18%")
        self.assertContains(r, "66.24")
        self.assertContains(r, "434.26")
        pdf = self.client.get(reverse("claim_repairs_pdf", args=[claim.pk]), {"company": co.pk})
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content[:5] == b"%PDF-")

    def test_loss_of_use_letter_pdf(self):
        from decimal import Decimal

        from .models import Company

        co = Company.objects.create(
            name="eCabs Ltd", address="St Julians, Malta",
            email="motorclaims@ecabs.com.mt", logo_static="img/companies/ecabs.png",
        )
        claim = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="BLY438",
            vehicle_make_model="Peugeot 208", third_party_registration="GCF078",
            tp_claim_number="M24109061", loe_days=3,
            loe_daily_rate=Decimal("159.58"), created_by=self.user,
        )
        # Documents page (loss-of-earnings summary)
        r = self.client.get(
            reverse("claim_documents", args=[claim.pk]),
            {"doc": "lossofuse", "company": co.pk},
        )
        self.assertContains(r, "Loss of use")
        self.assertContains(r, "3 days")
        # PDF
        r = self.client.get(reverse("claim_lou_pdf", args=[claim.pk]), {"company": co.pk})
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content[:5] == b"%PDF-")

    def test_invoice_pdf(self):
        from decimal import Decimal

        from .models import Company

        co = Company.objects.create(
            name="Vai Drive Co Ltd.", address="Paceville, Malta",
            bank_name="Banif Bank", logo_static="img/companies/vai.png",
        )
        claim = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-PDF",
            labour_amount=Decimal("100.00"), created_by=self.user,
        )
        r = self.client.get(
            reverse("claim_invoice_pdf", args=[claim.pk]), {"company": co.pk}
        )
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r["Content-Disposition"].startswith("inline"))
        self.assertTrue(r.content[:5] == b"%PDF-")

    def test_data_dashboard(self):
        Claim.objects.create(status=Claim.Status.OPEN, created_by=self.user)
        r = self.client.get(reverse("data_dashboard"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Claims by status")
        self.assertContains(r, "chartData")

    def test_bills_pending_by_insurer(self):
        from datetime import date
        from decimal import Decimal

        # Billed, still outstanding
        Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-PEND",
            third_party_insurer="Elmo Insurance",
            bills_sent_on=date(2026, 5, 1), parts_amount=Decimal("500.00"),
            created_by=self.user,
        )
        # Billed but fully paid -> not pending
        Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-PAID",
            third_party_insurer="Atlas",
            bills_sent_on=date(2026, 5, 1), parts_amount=Decimal("100.00"),
            amount_paid=Decimal("100.00"), created_by=self.user,
        )
        r = self.client.get(reverse("bills_report"))
        self.assertContains(r, "Bills pending")
        self.assertContains(r, "ECB-PEND")
        self.assertNotContains(r, "ECB-PAID")
        # Filter by a different insurer hides it
        r = self.client.get(reverse("bills_report"), {"pending_insurer": "GasanMamo"})
        self.assertNotContains(r, "ECB-PEND")

    def test_workflow_survey_liability_to_chasing(self):
        from datetime import date

        claim = Claim.objects.create(
            status=Claim.Status.AWAITING_SURVEY, created_by=self.user
        )
        # Not yet: survey received but liability not accepted
        claim.survey_in_hand = True
        claim.run_workflow()
        self.assertEqual(claim.status, Claim.Status.AWAITING_SURVEY)
        # Survey received + liability accepted -> chasing (awaiting insurer)
        claim.liability = Claim.Liability.ACCEPTED
        claim.run_workflow()
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.AWAITING_INSURER)
        self.assertIsNotNone(claim.chase_on)

    def test_workflow_closes_when_fully_paid(self):
        from decimal import Decimal

        claim = Claim.objects.create(
            status=Claim.Status.AWAITING_INSURER,
            parts_amount=Decimal("500.00"), created_by=self.user,
        )
        claim.run_workflow()  # outstanding 500 -> stays open
        self.assertEqual(claim.status, Claim.Status.AWAITING_INSURER)
        claim.amount_paid = Decimal("200.00")
        claim.run_workflow()  # still 300 outstanding -> stays open
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.AWAITING_INSURER)
        claim.amount_paid = Decimal("500.00")
        claim.run_workflow()  # fully paid -> closed
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.CLOSED)

    def test_full_automation_cycle(self):
        """Walk one claim through the whole lifecycle end-to-end."""
        from datetime import timedelta
        from decimal import Decimal

        from .services.auto_reminders import sync_auto_reminders

        today = timezone.localdate()

        def autos(text):
            return Reminder.objects.filter(
                is_auto=True, claim=claim, title__contains=text
            ).exists()

        # 1. Open the claim
        claim = Claim.objects.create(created_by=self.user)
        claim.submit()
        self.assertEqual(claim.status, Claim.Status.OPEN)

        # 2. Book a survey 2 days out -> day-before reminder
        claim.survey_booked = True
        claim.survey_date = today + timedelta(days=2)
        claim.save()
        sync_auto_reminders()
        self.assertTrue(autos("Survey tomorrow"))

        # 3. Survey date passes, still not received -> chase-survey reminder
        claim.survey_date = today - timedelta(days=1)
        claim.save()
        sync_auto_reminders()
        self.assertTrue(autos("Chase survey report"))
        self.assertFalse(autos("Survey tomorrow"))  # old one cleared

        # 4. Survey received but liability disputed -> monthly liability chase, still open
        claim.survey_in_hand = True
        claim.liability = Claim.Liability.DISPUTED
        claim.save()
        claim.run_workflow()
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.OPEN)
        sync_auto_reminders()
        self.assertTrue(autos("Chase liability"))

        # 5. Liability accepted -> claim becomes a chasing claim
        claim.liability = Claim.Liability.ACCEPTED
        claim.save()
        claim.run_workflow()
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.AWAITING_INSURER)
        self.assertIsNotNone(claim.chase_on)

        # 6. Bills sent with an amount outstanding -> monthly payment chase
        claim.bills_sent_on = today
        claim.parts_amount = Decimal("500.00")
        claim.insurer_contact = "MSI claims — 2123 4567"
        claim.save()
        sync_auto_reminders()
        self.assertTrue(autos("Chase payment"))

        # 7. Partial payment -> every cent still open
        claim.amount_paid = Decimal("200.00")
        claim.save()
        claim.run_workflow()
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.AWAITING_INSURER)

        # 8. Fully paid -> auto-closed
        claim.amount_paid = Decimal("500.00")
        claim.save()
        claim.run_workflow()
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.CLOSED)
        self.assertEqual(claim.outstanding_amount, Decimal("0.00"))

    def test_invoice_date_is_manual(self):
        import os
        import tempfile
        from datetime import date
        from decimal import Decimal

        import fitz

        from .models import Company

        co = Company.objects.create(
            name="Vai Drive Co Ltd.", address="Malta", logo_static="img/companies/vai.png"
        )
        claim = Claim.objects.create(
            status=Claim.Status.OPEN, bills_sent_on=date(2026, 5, 20),
            parts_amount=Decimal("100.00"), created_by=self.user,
        )
        # The documents page offers a manual date input and does NOT auto-fill
        # the invoice date from bills_sent_on.
        r = self.client.get(
            reverse("claim_documents", args=[claim.pk]),
            {"doc": "statement", "company": co.pk},
        )
        self.assertContains(r, "Invoice date")
        self.assertContains(r, 'name="date"')
        self.assertNotContains(r, "20/05/2026")

        def pdf_text(params):
            resp = self.client.get(reverse("claim_invoice_pdf", args=[claim.pk]), params)
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
                f.write(resp.content)
                path = f.name
            t = fitz.open(path).load_page(0).get_text().replace("\n", " ")
            os.unlink(path)
            return t

        # Typed date appears; no date -> dash, never today.
        self.assertIn("Invoice date: 07/06/2026", pdf_text({"company": co.pk, "date": "2026-06-07"}))
        self.assertIn("Invoice date: —", pdf_text({"company": co.pk}))

    def test_workflow_reminders(self):
        from datetime import timedelta

        from .services.auto_reminders import sync_auto_reminders

        today = timezone.localdate()
        # Survey booked in 3 days -> day-before reminder
        c1 = Claim.objects.create(
            status=Claim.Status.AWAITING_SURVEY, survey_booked=True,
            survey_date=today + timedelta(days=3), created_by=self.user,
        )
        # Liability disputed -> monthly chase
        c2 = Claim.objects.create(
            status=Claim.Status.OPEN, liability=Claim.Liability.DISPUTED,
            third_party_insurer="Elmo", created_by=self.user,
        )
        sync_auto_reminders()
        sync_auto_reminders()  # idempotent
        titles = list(Reminder.objects.filter(is_auto=True).values_list("title", flat=True))
        self.assertTrue(any("Survey tomorrow" in t for t in titles))
        self.assertTrue(any("Chase liability" in t for t in titles))
        self.assertEqual(
            Reminder.objects.filter(is_auto=True, title__contains="Survey tomorrow").count(), 1
        )

    def test_open_claims_estimate(self):
        from decimal import Decimal

        claim = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-EST", created_by=self.user
        )
        r = self.client.get(reverse("claim_group", args=["open"]))
        self.assertContains(r, "Total bill")
        # a claim with no real amounts shows the editable estimate cell
        self.assertContains(r, "estimate")
        self.assertFalse(claim.bill_is_actual)
        # Inline-save an estimate
        self.client.post(
            reverse("claim_set_estimate", args=[claim.pk]),
            {"estimate_amount": "1500.00", "next": reverse("claim_group", args=["open"])},
        )
        claim.refresh_from_db()
        self.assertEqual(claim.estimate_amount, Decimal("1500.00"))
        self.assertEqual(claim.effective_bill, Decimal("1500.00"))

    def test_custom_claim_status(self):
        from .models import Claim, ClaimStatus

        # add a custom status via the manage page
        self.client.post(reverse("claim_statuses_manage"),
                         {"action": "add", "name": "Awaiting parts"})
        cs = ClaimStatus.objects.get(name="Awaiting parts")
        self.assertTrue(cs.is_active)

        claim = Claim.objects.create(status=Claim.Status.OPEN,
                                     vehicle_registration="CSTT", created_by=self.user)
        # it shows in the status dropdown on the list
        self.assertContains(self.client.get(reverse("claim_group", args=["open"])),
                            "Awaiting parts")
        # a claim can be set to the custom status
        self.client.post(reverse("claim_set_status", args=[claim.pk]),
                         {"status": cs.slug, "next": reverse("claim_group", args=["open"])})
        claim.refresh_from_db()
        self.assertEqual(claim.status, cs.slug)
        self.assertEqual(claim.status_label, "Awaiting parts")

        # hiding removes it from the dropdown but keeps the label on the claim
        self.client.post(reverse("claim_statuses_manage"),
                         {"action": "toggle", "pk": cs.pk})
        cs.refresh_from_db()
        self.assertFalse(cs.is_active)
        self.assertEqual(claim.status_label, "Awaiting parts")

    def test_open_claims_bill_switches_to_actual(self):
        from decimal import Decimal

        claim = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-ACT",
            estimate_amount=Decimal("999.00"), labour_amount=Decimal("800.00"),
            amount_paid=Decimal("200.00"), created_by=self.user,
        )
        # real amounts entered → actual outstanding supersedes the estimate
        self.assertTrue(claim.bill_is_actual)
        self.assertEqual(claim.effective_bill, Decimal("600.00"))
        r = self.client.get(reverse("claim_group", args=["open"]))
        self.assertContains(r, "actual")

    def test_claim_sheets_sort(self):
        from datetime import date

        old = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-OLD",
            accident_date=date(2025, 1, 5), created_by=self.user,
        )
        new = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-NEW",
            accident_date=date(2026, 6, 20), created_by=self.user,
        )
        body = self.client.get(
            reverse("claim_group", args=["open"]), {"sort": "date_asc"}
        ).content.decode()
        self.assertLess(body.index("ECB-OLD"), body.index("ECB-NEW"))
        body = self.client.get(
            reverse("claim_group", args=["open"]), {"sort": "date_desc"}
        ).content.decode()
        self.assertLess(body.index("ECB-NEW"), body.index("ECB-OLD"))
        # Claims list accepts the sort param too.
        self.assertEqual(
            self.client.get(reverse("claim_list"), {"sort": "case"}).status_code, 200
        )

    def test_case_number_assigned_on_submit(self):
        a = Claim.objects.create()
        self.assertIsNone(a.case_number)
        self.assertEqual(a.case_ref, "")
        a.submit()
        a.refresh_from_db()
        self.assertEqual(a.case_number, 1)
        self.assertEqual(a.case_ref, "VD-00001")
        # Second claim gets the next running number; never resets.
        b = Claim.objects.create(status=Claim.Status.OPEN)
        self.assertEqual(b.case_number, 2)

    def test_master_sheet_grouped_by_month_and_summary(self):
        from datetime import date

        Claim.objects.create(
            status=Claim.Status.OPEN, accident_date=date(2026, 3, 15),
            created_by=self.user,
        )
        Claim.objects.create(
            status=Claim.Status.CLOSED, accident_date=date(2025, 11, 2),
            created_by=self.user,
        )
        response = self.client.get(reverse("master_sheet"))
        self.assertContains(response, "March 2026")
        self.assertContains(response, "November 2025")
        self.assertContains(response, "Claims by status")
        # Newest month appears before the older one in the page.
        body = response.content.decode()
        self.assertLess(body.index("March 2026"), body.index("November 2025"))

    def test_master_overdue_flag_filter(self):
        from datetime import timedelta

        today = timezone.localdate()
        overdue = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-OVR",
            chase_on=today - timedelta(days=2), created_by=self.user,
        )
        Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="ECB-OK",
            chase_on=today + timedelta(days=5), created_by=self.user,
        )
        response = self.client.get(reverse("master_sheet"), {"flag": "overdue"})
        self.assertContains(response, "ECB-OVR")
        self.assertNotContains(response, "ECB-OK")

    def test_vehicle_add_page(self):
        from decimal import Decimal

        response = self.client.get(reverse("vehicle_add"))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(
            reverse("vehicle_add"),
            {"registration": "ECB-NEW", "make_model": "Kia Picanto",
             "insurance_amount": "450.00", "licence_amount": "60.00"},
        )
        vehicle = Vehicle.objects.get(registration="ECB-NEW")
        self.assertRedirects(response, reverse("vehicle_edit", args=[vehicle.pk]))
        # The year's costs entered on the add form are saved and show up.
        cost = vehicle.costs.get()
        self.assertEqual(cost.insurance_amount, Decimal("450.00"))
        self.assertEqual(cost.licence_amount, Decimal("60.00"))

        # "Save & add another" keeps us on the add form.
        response = self.client.post(
            reverse("vehicle_add"),
            {"registration": "ECB-NEW2", "make_model": "Kia Rio", "add_another": "1"},
        )
        self.assertRedirects(response, reverse("vehicle_add"))
        self.assertTrue(Vehicle.objects.filter(registration="ECB-NEW2").exists())

    def test_vehicle_monthly_budget(self):
        """Costs split by month paid and by company for budgeting."""
        from datetime import date
        from decimal import Decimal

        from claims.models import Vehicle, VehicleAdditionalCost, VehicleCost
        from claims.views import _fleet_cost_budget

        v1 = Vehicle.objects.create(registration="BUD-1", owner="eCabs EN")
        VehicleCost.objects.create(
            vehicle=v1, year=2026, insurance_amount=Decimal("320.00"),
            licence_amount=Decimal("45.00"), pay_date=date(2026, 3, 1))
        v2 = Vehicle.objects.create(registration="BUD-2", owner="Fast Drop")
        c2 = VehicleCost.objects.create(
            vehicle=v2, year=2026, insurance_amount=Decimal("280.00"),
            licence_amount=Decimal("40.00"), pay_date=date(2026, 4, 5))
        VehicleAdditionalCost.objects.create(
            cost=c2, amount=Decimal("100.00"), date_added=date(2026, 5, 10))

        budget = _fleet_cost_budget()
        self.assertEqual(budget["companies"], ["eCabs EN", "Fast Drop"])
        # Grand total = 365 + 320 + 100
        self.assertEqual(budget["grand_total"], Decimal("785.00"))
        # Rows are most-recent-month first (May, Apr, Mar)
        self.assertEqual([r["label"] for r in budget["rows"]],
                         ["May 2026", "Apr 2026", "Mar 2026"])
        # The page renders the breakdown
        r = self.client.get(reverse("vehicle_list"))
        self.assertContains(r, "Monthly cost breakdown")

    def test_sales_invoice_client_autofill_and_totals(self):
        """ECABS sales invoice: client register auto-fills the bill-to block,
        VAT-inclusive lines back out Net/VAT, and the PDF renders."""
        from datetime import date
        from decimal import Decimal

        from claims.models import ClientInsurer, SalesInvoice, SalesInvoiceLine

        client_rec = ClientInsurer.objects.create(
            name="PWO S.p.A. Malta Branch",
            address="Focus House, First Floor,\nMalta",
            customer_no="CUST000445", vat_reg_no="MT 2099-1236",
            company_reg_no="OC 1562", default_payment_terms="Net 15 days")

        # New invoice, then choosing the client copies its details into the
        # (still editable) bill-to snapshot.
        self.client.post(reverse("sales_invoice_new"))
        inv = SalesInvoice.objects.get()
        self.client.post(reverse("sales_invoice_edit", args=[inv.pk]), {
            "client": client_rec.pk, "invoice_no": "PSIN01920245",
            "document_date": "2026-02-05", "due_date": "2026-02-20",
            "vat_rate": "18", "status": "sent",
        })
        inv.refresh_from_db()
        self.assertEqual(inv.bill_customer_no, "CUST000445")
        self.assertEqual(inv.bill_company_reg_no, "OC 1562")
        self.assertEqual(inv.payment_terms, "Net 15 days")
        self.assertEqual(inv.due_date, date(2026, 2, 20))

        # Two VAT-inclusive lines of 10.00 each → Total 20, Net 16.95, VAT 3.05.
        SalesInvoiceLine.objects.create(
            invoice=inv, description="Door", original_amount_incl_vat=Decimal("10"))
        SalesInvoiceLine.objects.create(
            invoice=inv, description="Door", original_amount_incl_vat=Decimal("10"))
        self.assertEqual(inv.total, Decimal("20"))
        self.assertEqual(inv.net, Decimal("16.95"))
        self.assertEqual(inv.vat_amount, Decimal("3.05"))

        r = self.client.get(reverse("sales_invoice_pdf", args=[inv.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF-"))

    def test_parts_receipt_standalone_and_claim_link(self):
        """Parts receipt: optional open-claim link, VAT-exclusive lines with VAT
        added at the bottom, and a PDF in the ECABS layout."""
        from datetime import date
        from decimal import Decimal

        from claims.models import PartsReceipt, PartsReceiptLine

        claim = Claim.objects.create(
            vehicle_registration="ECB101", status=Claim.Status.OPEN,
            submitted_at=timezone.now(), created_by=self.user)

        # New receipt pre-filled from the open claim.
        self.client.post(reverse("parts_receipt_new"), {"claim": claim.pk})
        rec = PartsReceipt.objects.get()
        self.assertEqual(rec.claim_id, claim.pk)
        self.assertEqual(rec.vehicle_reg, "ECB101")

        # One Save button adds a blank line at the same time as the header.
        self.client.post(reverse("parts_receipt_edit", args=[rec.pk]), {
            "receipt_no": "PR01001", "receipt_date": "2023-05-17",
            "supplier": "Michael Attard Ltd", "reference": "ecabshir",
            "vehicle_reg": "ECB101", "vat_rate": "18", "status": "sent",
            "add_line": "1",
        })
        rec.refresh_from_db()
        line1 = rec.lines.get()

        # Save the line's values via the same form, and add a second line.
        self.client.post(reverse("parts_receipt_edit", args=[rec.pk]), {
            "receipt_no": "PR01001", "receipt_date": "2023-05-17", "vat_rate": "18",
            f"line_{line1.pk}_part_code": "1618037980",
            f"line_{line1.pk}_description": "Bracket set",
            f"line_{line1.pk}_quantity": "1",
            f"line_{line1.pk}_unit_price": "25.95",
            "add_line": "1",
        })
        line1.refresh_from_db()
        self.assertEqual(line1.part_code, "1618037980")
        self.assertEqual(line1.unit_price, Decimal("25.95"))
        line2 = rec.lines.exclude(pk=line1.pk).get()
        PartsReceiptLine.objects.filter(pk=line2.pk).update(
            part_code="98120622", description="Grille",
            quantity=Decimal("2"), unit_price=Decimal("62.97"))

        rec.refresh_from_db()
        self.assertEqual(rec.receipt_date, date(2023, 5, 17))
        # Subtotal 25.95 + 2×62.97 = 151.89; VAT 27.34; total 179.23.
        self.assertEqual(rec.subtotal, Decimal("151.89"))
        self.assertEqual(rec.vat_amount, Decimal("27.34"))
        self.assertEqual(rec.total, Decimal("179.23"))

        r = self.client.get(reverse("parts_receipt_pdf", args=[rec.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF-"))

    def test_claim_can_hold_several_statuses(self):
        """Several statuses tick on one claim; a finishing one closes it."""
        claim = Claim.objects.create(
            vehicle_registration="ECB-900", status=Claim.Status.OPEN,
            submitted_at=timezone.now(), created_by=self.user)

        # Tick two at once from the open-claims table.
        self.client.post(reverse("claim_set_status", args=[claim.pk]), {
            "status": [Claim.Status.AWAITING_SURVEY, Claim.Status.AWAITING_INSURER],
            "next": reverse("claim_group", args=["open"]),
        })
        claim.refresh_from_db()
        self.assertEqual(claim.status, Claim.Status.AWAITING_SURVEY)
        self.assertEqual(
            set(claim.status_slugs),
            {Claim.Status.AWAITING_SURVEY, Claim.Status.AWAITING_INSURER},
        )
        self.assertFalse(claim.is_finished)

        # Both statuses find it, and it is still an open claim.
        for slug in (Claim.Status.AWAITING_SURVEY, Claim.Status.AWAITING_INSURER):
            self.assertContains(
                self.client.get(reverse("claim_list"), {"status": slug}),
                claim.reference)
        self.assertContains(
            self.client.get(reverse("claim_group", args=["open"])), claim.reference)

        # Ticking a finishing status closes it, even alongside an active one.
        self.client.post(reverse("claim_set_status", args=[claim.pk]), {
            "status": [Claim.Status.AWAITING_INSURER, Claim.Status.SETTLED],
            "next": reverse("claim_group", args=["open"]),
        })
        claim.refresh_from_db()
        self.assertTrue(claim.is_finished)
        self.assertNotContains(
            self.client.get(reverse("claim_group", args=["open"])), claim.reference)
        self.assertContains(
            self.client.get(reverse("claim_group", args=["closed"])), claim.reference)

        # Unticking it reopens the claim — no leftovers from the previous set.
        self.client.post(reverse("claim_set_status", args=[claim.pk]), {
            "status": [Claim.Status.OPEN],
            "next": reverse("claim_group", args=["open"]),
        })
        claim.refresh_from_db()
        self.assertEqual(claim.status_slugs, [Claim.Status.OPEN])
        self.assertFalse(claim.is_finished)

    def test_claims_master_import(self):
        """The master tracker imports into claims whose own arithmetic matches
        the sheet's total, re-imports without duplicating, and flags a row whose
        figures don't add up."""
        import io
        from decimal import Decimal

        import openpyxl

        from claims.services.claims_master_import import import_master_sheet

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([None, "25"])  # the stray cell above the real header
        ws.append([
            "Calendar Chaser", "Our Reg No", "TP Reg No", "Accident Date",
            "TP Claim No", "MSI Claim No", "TP Insurance", "RAR / ETARS / F2R",
            "Drivable?", "Survey / Photos", "Survey In Hand?",
            "Estimate Claim Amt", "Referred RGM / etc",
            "Claim Phase / Next Action", "Bills Sent", "Invoice No",
            "Labour (Net)", "Spray + Material (Net)", "Parts",
            "Loss of Earnings", "LOE (Days) - Daily Amt", "Others",
            "Settlement", "Amount Paid", "Offset", "TOTAL - Claim Amount",
        ])
        # Coherent: 243.6 + 429.53 + 800.21 + 0 + 0 − 1023.34 − 240 = 210.00
        ws.append([
            "2026-07-14", "RQZ031", "BUS013", "2023-08-14", "", "C34-277148",
            "Middlesea", "Etars", "YES", "", "Yes", "", "", "O/S payment",
            "2026-01-10", "100050", 243.6, 429.53, 800.21, 0, "", 0, 0,
            1023.34, 240, 210,
        ])
        # LOE split "4 days - € 185.62" = 742.48, plus 35 of Others.
        ws.append([
            "2026-06-30", "ELY236", "GLY062", "2023-10-30", "2023HAR2414", "",
            "Argus", "F 2 R", "NO - CLAIMING FROM 30/10", "", "", "", "", "",
            "", "500136", 100, 0, 0, 742.48, "4 days - € 185.62", 35, 0, 0, 0,
            877.48,
        ])
        # Paid in full: the settlement covers more than the itemised recovery,
        # so nothing is owed even though the sheet still shows a total.
        ws.append([
            "", "GQZ976", "ALH001", "2026-04-25", "M20262065", "C34-302964",
            "Atlas", "Etars", "", "", "", "", "", "", "", "", 0, 0, 0, 402.54,
            "", 0, 4500, 4500, 0, 402.54,
        ])
        ws.append([None] * 26)  # dragged-down blank row
        buf = io.BytesIO()
        wb.save(buf)

        buf.seek(0)
        dry = import_master_sheet(buf, commit=False)
        self.assertEqual(dry["rows"], 3)
        self.assertEqual(dry["created"], 3)
        self.assertEqual(Claim.objects.count(), 0)  # dry run writes nothing
        # Paid off, so it is reported as settled rather than as a bad sum.
        self.assertEqual(dry["warnings"], [])
        self.assertTrue(any("paid in full" in p["issue"] for p in dry["problems"]))

        buf.seek(0)
        report = import_master_sheet(buf, commit=True, user=self.user)
        self.assertEqual((report["created"], report["updated"]), (3, 0))
        self.assertEqual(Claim.objects.count(), 3)

        first = Claim.objects.get(insurer_claim_number="C34-277148")
        self.assertEqual(first.vehicle_registration, "RQZ031")
        self.assertEqual(first.report_type, Claim.ReportType.ETARS)
        self.assertTrue(first.drivable)
        self.assertEqual(first.outstanding_amount, Decimal("210.00"))

        # LOE is stored split, and "Others" rides on an OtherCharge so it counts.
        second = Claim.objects.get(tp_claim_number="2023HAR2414")
        self.assertEqual((second.loe_days, second.loe_daily_rate),
                         (4, Decimal("185.62")))
        self.assertEqual(second.loss_of_earnings, Decimal("742.48"))
        self.assertEqual(second.other_charges_total, Decimal("35.00"))
        self.assertEqual(second.outstanding_amount, Decimal("877.48"))

        # A paid-off claim is settled and owes nothing — never a negative bill.
        paid = Claim.objects.get(vehicle_registration="GQZ976")
        self.assertEqual(paid.status, Claim.Status.SETTLED)
        self.assertEqual(paid.outstanding_amount, Decimal("0"))
        self.assertTrue(paid.is_paid_in_full)

        # Re-importing the same sheet updates rather than duplicates.
        buf.seek(0)
        again = import_master_sheet(buf, commit=True, user=self.user)
        self.assertEqual((again["created"], again["updated"]), (0, 3))
        self.assertEqual(Claim.objects.count(), 3)
        self.assertEqual(OtherCharge.objects.filter(claim=second).count(), 1)

    def test_master_sheet_and_csv(self):
        from decimal import Decimal

        claim = Claim.objects.create(
            vehicle_registration="ECB-850",
            labour_amount=Decimal("100.00"),
            parts_amount=Decimal("50.00"),
            created_by=self.user,
        )
        response = self.client.get(reverse("master_sheet"))
        self.assertContains(response, claim.reference)
        self.assertContains(response, "150.00")  # computed total

        response = self.client.get(reverse("master_sheet_csv"))
        self.assertEqual(response["Content-Type"], "text/csv")
        body = response.content.decode()
        self.assertIn("TOTAL claim", body)
        self.assertIn(claim.reference, body)
        self.assertIn("150.00", body)

    def test_standalone_entry_forms(self):
        claim = Claim.objects.create(created_by=self.user)
        for kind in ("photo", "survey", "email", "billing"):
            response = self.client.get(reverse("record_add", args=[kind]))
            self.assertEqual(response.status_code, 200, kind)

        response = self.client.post(
            reverse("record_add", args=["survey"]),
            {"claim": claim.pk, "surveyor_name": "P. Attard", "status": "requested"},
        )
        self.assertRedirects(response, reverse("claim_tab", args=[claim.pk, "surveys"]))
        self.assertEqual(claim.surveys.count(), 1)

        response = self.client.post(
            reverse("record_add", args=["billing"]),
            {
                "claim": claim.pk,
                "description": "Towing",
                "category": "towing",
                "amount": "120.00",
                "status": "pending",
            },
        )
        self.assertRedirects(response, reverse("claim_tab", args=[claim.pk, "billing"]))
        self.assertEqual(claim.billing_items.count(), 1)

    def test_login_required(self):
        self.client.logout()
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response.url)


from django.core import mail  # noqa: E402
from django.test import override_settings  # noqa: E402


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class EmailSendTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("staff", password="pw")
        self.client.force_login(self.user)

    def test_send_email_delivers_and_logs(self):
        claim = Claim.objects.create(
            status=Claim.Status.OPEN, third_party_insurer="Mapfre",
            insurer_email="claims@mapfre.com.mt", created_by=self.user,
        )
        r = self.client.post(
            reverse("email_send", args=[claim.pk]),
            {"to": "claims@mapfre.com.mt", "subject": "Reminder",
             "body": "Please settle.", "attach": ""},
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["claims@mapfre.com.mt"])
        log = claim.emails.get()
        self.assertEqual(log.direction, "out")
        self.assertEqual(log.to_address, "claims@mapfre.com.mt")
        self.assertEqual(log.logged_by, self.user)

    def test_send_email_with_pdf_attachment(self):
        from decimal import Decimal

        from .models import Company

        co = Company.objects.create(
            name="eCabs Ltd", address="St Julians, Malta",
            email="motorclaims@ecabs.com.mt", logo_static="img/companies/ecabs.png",
        )
        claim = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="BLY438",
            tp_claim_number="M24109061", labour_amount=Decimal("200.00"),
            created_by=self.user,
        )
        r = self.client.post(
            reverse("email_send", args=[claim.pk]),
            {"company": co.pk, "to": "claims@mapfre.com.mt",
             "subject": "Repairs receipt", "body": "See attached.",
             "attach": "repairs"},
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(mail.outbox), 1)
        msg = mail.outbox[0]
        self.assertEqual(msg.from_email, "motorclaims@ecabs.com.mt")
        self.assertEqual(len(msg.attachments), 1)
        name, content, mimetype = msg.attachments[0]
        self.assertTrue(name.endswith(".pdf"))
        self.assertEqual(mimetype, "application/pdf")
        self.assertTrue(content[:5] == b"%PDF-")
        self.assertEqual(claim.emails.count(), 1)

    def test_attachment_requires_company(self):
        claim = Claim.objects.create(status=Claim.Status.OPEN, created_by=self.user)
        r = self.client.post(
            reverse("email_send", args=[claim.pk]),
            {"to": "x@y.com", "subject": "Hi", "body": "Body", "attach": "statement"},
        )
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Choose a company to attach a document.")
        self.assertEqual(len(mail.outbox), 0)
        self.assertEqual(claim.emails.count(), 0)

    def test_default_chase_message(self):
        from decimal import Decimal

        from .services.email import default_chase_message

        claim = Claim.objects.create(
            status=Claim.Status.OPEN, third_party_insurer="Elmo",
            parts_amount=Decimal("500.00"), created_by=self.user,
        )
        subject, body = default_chase_message(claim)
        self.assertIn("Elmo", body)
        self.assertIn("500.00", body)
        self.assertTrue(subject.startswith("Payment reminder"))


class PasswordAndBackupTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("staff", password="oldpw12345")

    def test_change_password_requires_login(self):
        r = self.client.get(reverse("password_change"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r.url)

    def test_change_password_works(self):
        self.client.force_login(self.user)
        r = self.client.get(reverse("password_change"))
        self.assertEqual(r.status_code, 200)
        r = self.client.post(reverse("password_change"), {
            "old_password": "oldpw12345",
            "new_password1": "brandNew99xy",
            "new_password2": "brandNew99xy",
        })
        self.assertRedirects(r, reverse("password_change_done"))
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password("brandNew99xy"))

    def test_password_reset_page_loads(self):
        r = self.client.get(reverse("password_reset"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Reset password")

    def test_password_reset_never_500s_without_email(self):
        # Submitting still redirects to the neutral done page even if sending
        # is not possible (SafePasswordResetView swallows send failures).
        r = self.client.post(reverse("password_reset"), {"email": "staff@example.com"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(r.url, reverse("password_reset_done"))

    def test_backup_db_command(self):
        import os
        import tempfile
        from pathlib import Path

        from django.conf import settings
        from django.core.management import call_command
        from django.test import override_settings

        with tempfile.TemporaryDirectory() as tmp:
            dbfile = Path(tmp) / "db.sqlite3"
            dbfile.write_bytes(b"SQLite format 3\x00sample")
            dbs = {"default": {"ENGINE": "django.db.backends.sqlite3", "NAME": str(dbfile)}}
            with override_settings(DATABASES=dbs, BASE_DIR=Path(tmp)):
                call_command("backup_db")
            made = list((Path(tmp) / "backups").glob("db-*.sqlite3"))
            self.assertEqual(len(made), 1)
            self.assertEqual(made[0].read_bytes(), dbfile.read_bytes())


class AuditTrailTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("auditor", password="pw")
        self.client.force_login(self.user)

    def test_audit_trail_shows_changes(self):
        from decimal import Decimal

        # A create and an update should both appear.
        claim = Claim.objects.create(
            status=Claim.Status.OPEN, vehicle_registration="AUD-1", created_by=self.user,
        )
        claim.parts_amount = Decimal("100.00")
        claim.save()
        Vehicle.objects.create(registration="AUD-VEH", make_model="Test")

        r = self.client.get(reverse("audit_trail"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Audit trail")
        self.assertContains(r, claim.reference)   # claim listed
        self.assertContains(r, "AUD-VEH")          # vehicle listed
        self.assertContains(r, "created")
        self.assertContains(r, "updated")
        self.assertContains(r, "parts_amount")     # changed field named

    def test_audit_trail_filter_by_kind(self):
        Claim.objects.create(status=Claim.Status.OPEN, vehicle_registration="AUD-2", created_by=self.user)
        Vehicle.objects.create(registration="ONLYVEH", make_model="Test")
        r = self.client.get(reverse("audit_trail"), {"kind": "vehicle"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "ONLYVEH")
        self.assertNotContains(r, "AUD-2")

    def test_audit_trail_requires_login(self):
        self.client.logout()
        r = self.client.get(reverse("audit_trail"))
        self.assertEqual(r.status_code, 302)


class ReminderPopupTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("staff", password="pw")
        self.client.force_login(self.user)

    def test_popup_shows_when_reminder_due(self):
        Reminder.objects.create(
            title="Chase Elmo payment",
            due_at=timezone.now() - timedelta(hours=1),  # overdue
        )
        r = self.client.get(reverse("home"))
        self.assertContains(r, "Reminders due")
        self.assertContains(r, "Chase Elmo payment")

    def test_no_popup_when_nothing_due(self):
        # A future reminder is not "due" yet.
        Reminder.objects.create(title="Later", due_at=timezone.now() + timedelta(days=3))
        r = self.client.get(reverse("home"))
        self.assertNotContains(r, "Reminders due")

    def test_completed_reminder_no_popup(self):
        Reminder.objects.create(
            title="Done one", due_at=timezone.now() - timedelta(hours=2),
            completed_at=timezone.now(),
        )
        r = self.client.get(reverse("home"))
        self.assertNotContains(r, "Reminders due")


class InvoiceDateAndNoteTests(TestCase):
    def _company(self, name):
        from .models import Company
        return Company.objects.create(name=name, address="Malta", logo_static="img/companies/vai.png")

    def test_invoice_date_blank_without_bill_date(self):
        from decimal import Decimal

        from .services.invoice_pdf import build_invoice_pdf
        co = self._company("Vai Drive Co Ltd.")
        claim = Claim.objects.create(status=Claim.Status.OPEN, parts_amount=Decimal("100.00"))
        # No bills_sent_on -> the PDF must not stamp today's date.
        import datetime, fitz, tempfile, os
        pdf = build_invoice_pdf(claim, co)
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            f.write(pdf); path = f.name
        text = fitz.open(path).load_page(0).get_text().replace("\n", " ")
        os.unlink(path)
        self.assertIn("Invoice date: —", text)
        self.assertNotIn(datetime.date.today().strftime("%d/%m/%Y"), text)

    def test_30pct_note_ecabs_only(self):
        from .services.invoice_pdf import build_lou_pdf
        import fitz, tempfile, os
        vai = self._company("Vai Drive Co Ltd.")
        ecabs = self._company("eCabs Ltd")
        from decimal import Decimal
        claim = Claim.objects.create(status=Claim.Status.OPEN, loe_days=8, loe_daily_rate=Decimal("50.00"))

        def note_present(company):
            pdf = build_lou_pdf(claim, company)
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
                f.write(pdf); path = f.name
            text = fitz.open(path).load_page(0).get_text()
            os.unlink(path)
            return "30% running expenses" in text

        self.assertTrue(note_present(ecabs))
        self.assertFalse(note_present(vai))


class OverviewPageTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("staff", password="pw")
        self.client.force_login(self.user)

    def test_overview_loads_with_companies(self):
        from .models import Company
        Company.objects.create(name="eCabs Ltd", logo_static="img/companies/ecabs.png")
        r = self.client.get(reverse("overview"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "The claims ecosystem")
        self.assertContains(r, "The claim lifecycle")
        self.assertContains(r, "eCabs Ltd")

    def test_overview_requires_login(self):
        self.client.logout()
        self.assertEqual(self.client.get(reverse("overview")).status_code, 302)


class ClaimsAtFaultTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("staff", password="pw")
        self.client.force_login(self.user)

    def test_only_at_fault_claims_listed(self):
        ours = Claim.objects.create(status=Claim.Status.OPEN, fault=Claim.Fault.OUR_DRIVER,
                                    vehicle_registration="OI-1", created_by=self.user)
        tp = Claim.objects.create(status=Claim.Status.OPEN, fault=Claim.Fault.THIRD_PARTY,
                                  vehicle_registration="TP-9", created_by=self.user)
        r = self.client.get(reverse("claims_at_fault"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "OI-1")
        self.assertNotContains(r, "TP-9")

    def test_inline_save_and_close(self):
        from decimal import Decimal
        c = Claim.objects.create(status=Claim.Status.OPEN, fault=Claim.Fault.OUR_DRIVER,
                                 created_by=self.user)
        # inline save of details / awaiting / estimate
        self.client.post(reverse("at_fault_update", args=[c.pk]),
                         {"details": "Rear-ended", "awaiting_from": "Mario",
                          "estimate_amount": "1250.50"})
        c.refresh_from_db()
        self.assertEqual(c.details, "Rear-ended")
        self.assertEqual(c.awaiting_from, "Mario")
        self.assertEqual(c.estimate_amount, Decimal("1250.50"))
        # close by decision
        self.client.post(reverse("at_fault_update", args=[c.pk]), {"action": "close"})
        c.refresh_from_db()
        self.assertEqual(c.status, Claim.Status.CLOSED)

    def test_at_fault_not_auto_closed_at_zero(self):
        # An at-fault claim with no outstanding must NOT auto-close.
        c = Claim.objects.create(status=Claim.Status.OPEN, fault=Claim.Fault.OUR_DRIVER,
                                 created_by=self.user)
        c.run_workflow()
        c.refresh_from_db()
        self.assertEqual(c.status, Claim.Status.OPEN)


class RepairTypeTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("staff", password="pw")
        self.client.force_login(self.user)

    def test_add_type_then_itemise_lines_with_vat(self):
        from decimal import Decimal

        from .models import Company, RepairLine, RepairType

        co = Company.objects.create(name="eCabs Ltd", address="Malta",
                                    logo_static="img/companies/ecabs.png")
        claim = Claim.objects.create(status=Claim.Status.OPEN, created_by=self.user)

        # add a new type — maintained list, not automated, no auto-select
        self.client.post(reverse("repair_type_add", args=[claim.pk]),
                         {"name": "Windscreen replacement"})
        rt = RepairType.objects.get(name="Windscreen replacement")
        self.assertFalse(claim.repair_lines.exists())

        body = RepairType.objects.get(name="Bodywork")

        # tick two types with per-line net costs
        self.client.post(reverse("claim_set_repair_lines", args=[claim.pk]), {
            f"line_{rt.pk}": "on", f"cost_{rt.pk}": "150.00",
            f"line_{body.pk}": "on", f"cost_{body.pk}": "50.00",
        })
        claim.refresh_from_db()
        self.assertEqual(claim.repair_lines.count(), 2)
        # net total is the sum of the ticked lines
        self.assertEqual(claim.repair_lines_total, Decimal("200.00"))
        self.assertEqual(claim.repairs_receipt_net, Decimal("200.00"))

        # unticking removes a line, keeps the other
        self.client.post(reverse("claim_set_repair_lines", args=[claim.pk]), {
            f"line_{rt.pk}": "on", f"cost_{rt.pk}": "150.00",
        })
        claim.refresh_from_db()
        self.assertEqual(claim.repair_lines.count(), 1)
        self.assertEqual(claim.repairs_receipt_net, Decimal("150.00"))

        # receipt PDF renders the itemised lines
        r = self.client.get(reverse("claim_repairs_pdf", args=[claim.pk]), {"company": co.pk})
        self.assertEqual(r["Content-Type"], "application/pdf")

    def test_starter_types_seeded(self):
        from .models import RepairType
        self.assertTrue(RepairType.objects.filter(name="Bodywork").exists())

    def test_repair_line_selector_on_documents_page(self):
        claim = Claim.objects.create(status=Claim.Status.OPEN, created_by=self.user)
        r = self.client.get(reverse("claim_documents", args=[claim.pk]), {"doc": "repairs"})
        self.assertContains(r, "Types of repair &amp; cost")
        self.assertContains(r, "Add a new repair type")


class GarageJobTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("garage", password="pw")
        self.client.force_login(self.user)

    def test_seeded_rows_show_on_page(self):
        r = self.client.get(reverse("garage_jobs"))
        self.assertContains(r, "ACR Garage")
        self.assertContains(r, "Mohammed Asif Khaled")

    def test_add_edit_close_reopen_delete_row(self):
        from datetime import date

        from .models import GarageJob

        before = GarageJob.objects.count()
        self.client.post(reverse("garage_job_add"))
        job = GarageJob.objects.order_by("-id").first()
        self.assertEqual(GarageJob.objects.count(), before + 1)

        # inline edit
        self.client.post(reverse("garage_job_update", args=[job.pk]), {
            "accident_date": "2026-06-01", "survey_date": "2026-06-05",
            "plate_no": "ZZZ999", "client": "Test Client", "make": "Toyota",
            "claim_no": "CL123", "surveyor": "Jo", "insurance": "Elmo",
            "report_received": "2026-06-10", "go_ahead": "Yes, go",
        })
        job.refresh_from_db()
        self.assertEqual(job.plate_no, "ZZZ999")
        self.assertEqual(job.accident_date, date(2026, 6, 1))
        self.assertTrue(job.go_ahead_yes)

        # close then reopen
        self.client.post(reverse("garage_job_update", args=[job.pk]), {"action": "close"})
        job.refresh_from_db(); self.assertTrue(job.is_closed)
        self.client.post(reverse("garage_job_update", args=[job.pk]), {"action": "reopen"})
        job.refresh_from_db(); self.assertFalse(job.is_closed)

        # delete
        self.client.post(reverse("garage_job_update", args=[job.pk]), {"action": "delete"})
        self.assertFalse(GarageJob.objects.filter(pk=job.pk).exists())

    def test_job_items_and_labour_feed_invoice(self):
        from decimal import Decimal

        from .models import GarageInvoice, GarageJob, RepairType

        spray = RepairType.objects.filter(name__icontains="Spray").first()
        # Items (existing + new type) and labour are all submitted in one form.
        self.client.post(reverse("garage_job_new"), {
            "plate_no": "ITM9", "client": "FASTDROP", "make": "Toyota",
            "item_type": [str(spray.pk), ""], "item_new": ["", "Panel"],
            "item_price": ["120.00", "80.00"],
            "labour_desc": ["Panel beating"], "labour_hours": ["3"],
            "labour_rate": ["30"],
        })
        job = GarageJob.objects.get(plate_no="ITM9")
        self.assertEqual(job.items.count(), 2)
        self.assertEqual(job.items_total, Decimal("200.00"))
        self.assertEqual(job.labour.count(), 1)
        self.assertEqual(job.labour_total, Decimal("90.00"))  # 3 × 30
        self.assertEqual(job.effective_total, Decimal("290.00"))
        self.assertTrue(RepairType.objects.filter(name="Panel").exists())

        # → invoice makes one line per item AND one per labour line
        self.client.post(reverse("garage_job_to_invoice", args=[job.pk]))
        inv = GarageInvoice.objects.filter(created_by=self.user).order_by("-id").first()
        self.assertEqual(inv.lines.count(), 3)
        self.assertEqual(inv.subtotal, Decimal("290.00"))

        # editing the form replaces items/labour from the submitted rows
        self.client.post(reverse("garage_job_form", args=[job.pk]), {
            "plate_no": "ITM9", "item_type": [str(spray.pk)], "item_new": [""],
            "item_price": ["150.00"],
        })
        job.refresh_from_db()
        self.assertEqual(job.items.count(), 1)
        self.assertEqual(job.labour.count(), 0)
        self.assertEqual(job.effective_total, Decimal("150.00"))

    def test_add_and_edit_via_full_form(self):
        from datetime import date

        from .models import GarageJob

        # GET the blank form
        self.assertEqual(self.client.get(reverse("garage_job_new")).status_code, 200)
        # POST creates a job
        self.client.post(reverse("garage_job_new"), {
            "accident_date": "2026-07-01", "survey_date": "2026-07-03",
            "plate_no": "FRM123", "client": "Test Co", "make": "Toyota",
            "claim_no": "CL9", "surveyor": "Jo", "insurance": "Elmo",
            "report_received": "2026-07-05", "go_ahead": "Yes", "total": "500.00",
        })
        job = GarageJob.objects.get(plate_no="FRM123")
        self.assertEqual(job.accident_date, date(2026, 7, 1))
        self.assertEqual(str(job.total), "500.00")
        self.assertFalse(job.is_closed)
        # edit via the same form, mark closed
        self.client.post(reverse("garage_job_form", args=[job.pk]),
                         {"plate_no": "FRM123", "client": "Renamed", "is_closed": "on"})
        job.refresh_from_db()
        self.assertEqual(job.client, "Renamed")
        self.assertTrue(job.is_closed)


class GarageAccessTests(TestCase):
    """A garage-only user (Mario) is confined to the ACR Garage page; an admin
    (Frances) keeps full access and sees the same rows."""

    def setUp(self):
        from django.contrib.auth.models import Group

        User = get_user_model()
        self.mario = User.objects.create_user("mario", password="pw")
        group, _ = Group.objects.get_or_create(name="Garage")
        self.mario.groups.add(group)
        self.francis = User.objects.create_superuser("francis", password="pw")

    def test_make_garage_user_command(self):
        from django.core.management import call_command

        from .roles import is_garage_user
        call_command("make_garage_user", "luigi", password="pw", full_name="Luigi Verde")
        u = get_user_model().objects.get(username="luigi")
        self.assertTrue(is_garage_user(u))
        self.assertFalse(u.is_staff)
        self.assertEqual(u.first_name, "Luigi")

    def test_garage_user_confined_to_garage_page(self):
        self.client.force_login(self.mario)
        self.assertEqual(self.client.get(reverse("garage_jobs")).status_code, 200)
        # anything else bounces to the worklist
        for name in ("dashboard", "claim_list", "claims_at_fault", "master_sheet"):
            r = self.client.get(reverse(name))
            self.assertRedirects(r, reverse("garage_jobs"),
                                 fetch_redirect_response=False)

    def test_garage_user_nav_is_trimmed(self):
        self.client.force_login(self.mario)
        html = self.client.get(reverse("garage_jobs")).content
        self.assertIn(b"ACR Garage", html)
        self.assertNotIn(b"Dashboard", html)
        self.assertNotIn(b"Master sheet", html)

    def test_admin_has_full_access_and_sees_garage(self):
        self.client.force_login(self.francis)
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 200)
        html = self.client.get(reverse("garage_jobs")).content
        self.assertIn(b"Dashboard", html)  # full nav
        self.assertIn(b"ACR Garage", html)

    def test_edits_are_attributed_and_shared(self):
        from .models import GarageJob

        self.client.force_login(self.mario)
        self.client.post(reverse("garage_job_add"))
        job = GarageJob.objects.order_by("-id").first()
        self.client.post(reverse("garage_job_update", args=[job.pk]),
                         {"plate_no": "MAR123", "total": "1250.50", "go_ahead": ""})
        job.refresh_from_db()
        self.assertEqual(job.updated_by, self.mario)
        self.assertEqual(str(job.total), "1250.50")
        # Frances sees Mario's row and its total on the shared page
        self.client.force_login(self.francis)
        html = self.client.get(reverse("garage_jobs")).content
        self.assertIn(b"MAR123", html)
        self.assertIn(b"1250.50", html)


class GarageInvoiceTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import Group

        User = get_user_model()
        self.mario = User.objects.create_user("mario", password="pw")
        group, _ = Group.objects.get_or_create(name="Garage")
        self.mario.groups.add(group)
        self.client.force_login(self.mario)

    def _make_invoice(self):
        from .models import GarageInvoice

        self.client.post(reverse("garage_invoice_new"))
        return GarageInvoice.objects.order_by("-id").first()

    def test_totals_compute_from_lines(self):
        from decimal import Decimal

        from .models import GarageInvoice

        inv = self._make_invoice()
        self.client.post(reverse("garage_invoice_edit", args=[inv.pk]),
                         {"discount_pct": "10", "vat_rate": "18", "bill_to": "FASTDROP"})
        # two lines
        self.client.post(reverse("garage_invoice_line_add", args=[inv.pk]))
        self.client.post(reverse("garage_invoice_line_add", args=[inv.pk]))
        lines = list(inv.lines.all())
        self.client.post(reverse("garage_invoice_line_update", args=[inv.pk, lines[0].pk]),
                         {"reg_no": "HCF506", "description": "Wing", "amount": "100.00"})
        self.client.post(reverse("garage_invoice_line_update", args=[inv.pk, lines[1].pk]),
                         {"reg_no": "CEB981", "description": "Door", "amount": "100.00"})
        inv = GarageInvoice.objects.get(pk=inv.pk)
        self.assertEqual(inv.subtotal, Decimal("200.00"))
        self.assertEqual(inv.discount_amount, Decimal("20.00"))
        self.assertEqual(inv.subtotal_less_discount, Decimal("180.00"))
        self.assertEqual(inv.vat_amount, Decimal("32.40"))
        self.assertEqual(inv.balance_due, Decimal("212.40"))

    def test_defaults_are_acr_garage(self):
        inv = self._make_invoice()
        self.assertEqual(inv.issuer_name, "ACR Garage")
        self.assertEqual(inv.payable_to, "Mario Galea")
        self.assertTrue(inv.invoice_no.startswith("INV NO."))

    def test_pdf_renders(self):
        inv = self._make_invoice()
        self.client.post(reverse("garage_invoice_line_add", args=[inv.pk]))
        r = self.client.get(reverse("garage_invoice_pdf", args=[inv.pk]))
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertGreater(len(r.content), 800)

    def test_line_delete(self):
        inv = self._make_invoice()
        self.client.post(reverse("garage_invoice_line_add", args=[inv.pk]))
        line = inv.lines.first()
        self.client.post(reverse("garage_invoice_line_update", args=[inv.pk, line.pk]),
                         {"action": "delete"})
        self.assertEqual(inv.lines.count(), 0)


class GarageInvoicePrivacyTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import Group

        User = get_user_model()
        self.mario = User.objects.create_user("mario", password="pw")
        grp, _ = Group.objects.get_or_create(name="Garage")
        self.mario.groups.add(grp)
        self.francis = User.objects.create_superuser("francis", password="pw")

    def _new_invoice_as(self, user):
        from .models import GarageInvoice

        self.client.force_login(user)
        self.client.post(reverse("garage_invoice_new"))
        return GarageInvoice.objects.filter(created_by=user).order_by("-id").first()

    def test_garage_user_only_sees_own(self):
        mine = self._new_invoice_as(self.mario)
        theirs = self._new_invoice_as(self.francis)
        self.client.force_login(self.mario)
        html = self.client.get(reverse("garage_invoices")).content.decode()
        self.assertIn(mine.invoice_no, html)
        self.assertNotIn(theirs.invoice_no, html)

    def test_garage_user_cannot_open_others_invoice(self):
        theirs = self._new_invoice_as(self.francis)
        self.client.force_login(self.mario)
        r = self.client.get(reverse("garage_invoice_edit", args=[theirs.pk]))
        self.assertRedirects(r, reverse("garage_invoices"))
        r2 = self.client.get(reverse("garage_invoice_pdf", args=[theirs.pk]))
        self.assertRedirects(r2, reverse("garage_invoices"))

    def test_admin_defaults_to_vai_drive_but_can_filter_personal(self):
        from .models import GarageInvoice

        # Mario raises one Vai-Drive invoice and one personal one.
        group_inv = self._new_invoice_as(self.mario)
        personal_inv = self._new_invoice_as(self.mario)
        personal_inv.for_kind = GarageInvoice.ForKind.PERSONAL
        personal_inv.save()

        self.client.force_login(self.francis)
        # default shows Vai-Drive invoices (across owners), hides personal ones
        default = self.client.get(reverse("garage_invoices")).content.decode()
        self.assertIn(group_inv.invoice_no, default)
        self.assertNotIn(personal_inv.invoice_no, default)
        # For=personal lets her glance at the garage's personal invoices
        personal = self.client.get(reverse("garage_invoices"), {"for": "personal"}).content.decode()
        self.assertIn(personal_inv.invoice_no, personal)

    def test_status_filter(self):
        inv = self._new_invoice_as(self.francis)
        inv.status = "paid"; inv.save()
        self.client.force_login(self.francis)
        r = self.client.get(reverse("garage_invoices"), {"owner": "all", "status": "paid"})
        self.assertContains(r, inv.invoice_no)
        r2 = self.client.get(reverse("garage_invoices"), {"owner": "all", "status": "draft"})
        self.assertNotContains(r2, inv.invoice_no)

    def test_invoice_numbers_are_never_reused(self):
        """Deleting an invoice must not hand its number to the next one."""
        from django.db import IntegrityError

        from .models import GarageInvoice

        self.client.force_login(self.francis)
        numbers = []
        for _ in range(3):
            self.client.post(reverse("garage_invoice_new"))
            numbers.append(GarageInvoice.objects.order_by("-id").first().invoice_no)
        self.assertEqual(len(set(numbers)), 3)  # all different

        # Delete the newest, make another — the number must move on, not repeat.
        newest = GarageInvoice.objects.order_by("-id").first()
        self.client.post(reverse("garage_invoice_delete", args=[newest.pk]))
        self.client.post(reverse("garage_invoice_new"))
        self.assertNotIn(
            GarageInvoice.objects.order_by("-id").first().invoice_no, numbers)

        # Even wiping every invoice must not wind the counter back.
        highest = GarageInvoice.objects.order_by("-id").first().invoice_no
        GarageInvoice.objects.all().delete()
        self.client.post(reverse("garage_invoice_new"))
        self.assertNotEqual(
            GarageInvoice.objects.order_by("-id").first().invoice_no, highest)

        # And the database refuses a duplicate outright.
        taken = GarageInvoice.objects.order_by("-id").first().invoice_no
        with self.assertRaises(IntegrityError):
            GarageInvoice.objects.create(invoice_no=taken)

    def test_invoice_delete(self):
        """A garage invoice can be deleted, and its lines go with it."""
        from decimal import Decimal

        from .models import GarageInvoice, GarageInvoiceLine

        inv = self._new_invoice_as(self.francis)
        GarageInvoiceLine.objects.create(invoice=inv, description="Labour",
                                         amount=Decimal("100"), order=0)
        self.client.force_login(self.francis)
        # GET must not delete — only POST.
        self.assertEqual(
            self.client.get(reverse("garage_invoice_delete", args=[inv.pk])).status_code,
            405)
        self.assertTrue(GarageInvoice.objects.filter(pk=inv.pk).exists())

        r = self.client.post(reverse("garage_invoice_delete", args=[inv.pk]))
        self.assertRedirects(r, reverse("garage_invoices"))
        self.assertFalse(GarageInvoice.objects.filter(pk=inv.pk).exists())
        self.assertEqual(GarageInvoiceLine.objects.count(), 0)

    def test_job_to_invoice_prefills(self):
        from decimal import Decimal

        from .models import GarageInvoice, GarageJob

        job = GarageJob.objects.create(garage=GarageJob.Garage.ACL, client="FASTDROP",
                                       plate_no="GLY555", make="Toyota",
                                       total=Decimal("300.00"))
        self.client.force_login(self.mario)
        self.client.post(reverse("garage_job_to_invoice", args=[job.pk]))
        inv = GarageInvoice.objects.filter(created_by=self.mario).order_by("-id").first()
        self.assertEqual(inv.bill_to, "FASTDROP")
        self.assertEqual(inv.vehicle_reg, "GLY555")  # plate now sits in the header
        line = inv.lines.first()
        self.assertEqual(str(line.amount), "300.00")

    def test_metrics_page(self):
        self._new_invoice_as(self.francis)
        self.client.force_login(self.francis)
        r = self.client.get(reverse("garage_metrics"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Invoiced (total)")

    def test_worklist_search(self):
        from .models import GarageJob

        GarageJob.objects.create(garage=GarageJob.Garage.ACL, plate_no="ZED111",
                                 client="Findme Ltd")
        self.client.force_login(self.francis)
        r = self.client.get(reverse("garage_jobs"), {"q": "Findme"})
        self.assertContains(r, "ZED111")
        r2 = self.client.get(reverse("garage_jobs"), {"q": "Nothinghere"})
        self.assertNotContains(r2, "ZED111")


class MetricsSectionTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser("boss", password="pw")

    def test_metrics_page_renders_for_admin(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("metrics"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Total claims")
        self.assertContains(r, "ACR Garage")

    def test_garage_user_blocked_from_metrics(self):
        from django.contrib.auth.models import Group

        m = get_user_model().objects.create_user("g", password="pw")
        grp, _ = Group.objects.get_or_create(name="Garage")
        m.groups.add(grp)
        self.client.force_login(m)
        r = self.client.get(reverse("metrics"))
        self.assertRedirects(r, reverse("garage_jobs"), fetch_redirect_response=False)


class GeneralClaimsTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("gcuser", password="pw")
        self.client.force_login(self.user)

    def test_add_edit_close_and_search(self):
        from datetime import date

        from .models import GeneralClaim

        self.assertEqual(self.client.get(reverse("general_claim_new")).status_code, 200)
        self.client.post(reverse("general_claim_new"), {
            "claim_date": "2026-07-01", "brand": "Firefly", "our_reg": "FRF001",
            "driver_name": "Joe", "insurer": "Elmo", "claim_no": "GC1",
            "fault": "TP", "details": "rear hit", "status_note": "Awaiting RGM",
            "estimate_amount": "450.00",
        })
        row = GeneralClaim.objects.get(our_reg="FRF001")
        self.assertEqual(row.brand, "Firefly")
        self.assertEqual(row.fault, "TP")
        self.assertEqual(row.claim_date, date(2026, 7, 1))
        self.assertEqual(str(row.estimate_amount), "450.00")
        self.assertFalse(row.is_closed)

        # appears in Open section, search works
        r = self.client.get(reverse("general_claims"), {"q": "FRF001"})
        self.assertContains(r, "FRF001")
        self.assertNotContains(self.client.get(reverse("general_claims"), {"q": "zzz"}), "FRF001")

        # close from the list
        self.client.post(reverse("general_claim_toggle", args=[row.pk]), {"action": "close"})
        row.refresh_from_db()
        self.assertTrue(row.is_closed)

        # edit then delete
        self.client.post(reverse("general_claim_edit", args=[row.pk]),
                         {"our_reg": "FRF001", "details": "updated"})
        row.refresh_from_db()
        self.assertEqual(row.details, "updated")
        self.client.post(reverse("general_claim_delete", args=[row.pk]))
        self.assertFalse(GeneralClaim.objects.filter(pk=row.pk).exists())

    def test_garage_user_cannot_reach_general_claims(self):
        from django.contrib.auth.models import Group

        m = get_user_model().objects.create_user("gg", password="pw")
        grp, _ = Group.objects.get_or_create(name="Garage")
        m.groups.add(grp)
        self.client.force_login(m)
        r = self.client.get(reverse("general_claims"))
        self.assertRedirects(r, reverse("garage_jobs"), fetch_redirect_response=False)


class AccidentListTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("fran", password="pw")
        self.client.force_login(self.user)

    def _xlsx(self):
        import openpyxl
        from io import BytesIO

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Date of Acc", "Our Reg", "Driver Name", "TP Reg", "Vehicle Make",
                   "TP Driver Name", "Contact Details", "TP Owner Name",
                   "TP Owner Contact No", "TP Insurance", "Other Tps", "TP2 Owner Name",
                   "TP2 Contact No", "TP2 Insurance", "Report Type", "Fault (OI / TP)"])
        ws.append(["2021-01-01", "CLY142", "Isaac Zammit", "BBX937", "Nissan Qashqai",
                   "Michael Calleja", "99432967", "Michael Calleja", "99432967",
                   "MSI", "", "", "", "", "F 2 R", "TP"])
        ws.append(["2021-02-03", "FLY098", "Samwel Sammut", "JBH412", "Peugeot 206",
                   "Brian Vella", "79707569", "Brian Vella", "79707569", "MSI",
                   "", "", "", "", "Etars", "OI"])
        ws.append([None] * 16)  # blank row
        buf = BytesIO(); wb.save(buf); buf.seek(0)
        buf.name = "accidents.xlsx"
        return buf

    def test_upload_imports_records(self):
        from .models import AccidentRecord

        r = self.client.post(reverse("accident_list_upload"), {"file": self._xlsx()})
        self.assertRedirects(r, reverse("accident_list"))
        self.assertEqual(AccidentRecord.objects.count(), 2)
        rec = AccidentRecord.objects.get(our_reg="CLY142")
        self.assertEqual(rec.driver_name, "Isaac Zammit")
        self.assertEqual(rec.fault, "TP")
        self.assertEqual(str(rec.date_of_acc), "2021-01-01")

    def test_replace_clears_first(self):
        from .models import AccidentRecord

        AccidentRecord.objects.create(our_reg="OLD001")
        self.client.post(reverse("accident_list_upload"),
                         {"file": self._xlsx(), "replace": "on"})
        self.assertFalse(AccidentRecord.objects.filter(our_reg="OLD001").exists())
        self.assertEqual(AccidentRecord.objects.count(), 2)

    def test_search_and_manual_add(self):
        from .models import AccidentRecord

        self.client.post(reverse("accident_record_new"),
                         {"our_reg": "ZZZ999", "driver_name": "Jane", "fault": "OI"})
        rec = AccidentRecord.objects.get(our_reg="ZZZ999")
        self.assertEqual(rec.driver_name, "Jane")
        r = self.client.get(reverse("accident_list"), {"q": "ZZZ999"})
        self.assertContains(r, "ZZZ999")

    def test_admin_can_view_as_mario_and_switch_back(self):
        from django.contrib.auth.models import Group

        User = get_user_model()
        francis = User.objects.create_superuser("francis_imp", password="pw")
        mario = User.objects.create_user("mario_imp", password="pw",
                                         first_name="Mario")
        grp, _ = Group.objects.get_or_create(name="Garage")
        mario.groups.add(grp)

        self.client.force_login(francis)
        # button visible to admin
        self.assertContains(self.client.get(reverse("home")), "View Mario")
        # start impersonation → lands on the garage worklist, confined + banner
        self.client.get(reverse("impersonate_garage"))
        wl = self.client.get(reverse("garage_jobs"))
        self.assertContains(wl, "Viewing as")
        self.assertContains(wl, "My worklist")  # garage-user nav
        # confined like Mario
        self.assertRedirects(self.client.get(reverse("dashboard")),
                             reverse("garage_jobs"), fetch_redirect_response=False)
        # switch back
        self.client.get(reverse("stop_impersonate"))
        self.assertContains(self.client.get(reverse("home")), "Dashboard")

    def test_non_admin_cannot_impersonate(self):
        from django.contrib.auth.models import Group

        m = get_user_model().objects.create_user("m_imp2", password="pw")
        grp, _ = Group.objects.get_or_create(name="Garage")
        m.groups.add(grp)
        self.client.force_login(m)
        # garage user hitting the impersonate URL is just bounced to worklist
        self.assertRedirects(self.client.get(reverse("impersonate_garage")),
                             reverse("garage_jobs"), fetch_redirect_response=False)

    def test_garage_user_cannot_reach_accidents(self):
        from django.contrib.auth.models import Group

        mario = get_user_model().objects.create_user("mario_a", password="pw")
        grp, _ = Group.objects.get_or_create(name="Garage")
        mario.groups.add(grp)
        self.client.force_login(mario)
        r = self.client.get(reverse("accident_list"))
        self.assertRedirects(r, reverse("garage_jobs"), fetch_redirect_response=False)


class ImportAccidentsTests(TestCase):
    def _make_xlsx(self, path):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Date of Acc", "Our Reg", "Driver Name", "TP Reg", "Vehicle Make",
                   "TP Insurance", "Report Type", "Fault (OI / TP)"])
        import datetime
        ws.append([datetime.datetime(2021, 1, 1), "CLY142", "Isaac Zammit", "BBX937",
                   "Nissan Qashqai", "MSI", "F 2 R", "TP"])
        ws.append([datetime.datetime(2021, 2, 3), "FLY098", "Sam", "JBH412",
                   "Peugeot 206", "MSI", "Etars - 307250", "OI"])
        ws.append([None, "", "", "", "", "", "", ""])  # empty -> skipped
        wb.save(path)

    def test_import_maps_and_creates(self):
        import tempfile
        from django.core.management import call_command
        from .models import Claim

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        self._make_xlsx(path)
        call_command("import_accidents", path, status="closed")

        self.assertEqual(Claim.objects.count(), 2)  # empty row skipped
        a = Claim.objects.get(vehicle_registration="CLY142")
        self.assertEqual(a.driver_name, "Isaac Zammit")
        self.assertEqual(a.third_party_registration, "BBX937")
        self.assertEqual(a.third_party_insurer, "MSI")
        self.assertEqual(a.fault, Claim.Fault.THIRD_PARTY)
        self.assertEqual(a.report_type, Claim.ReportType.F2R)
        self.assertEqual(a.status, Claim.Status.CLOSED)
        b = Claim.objects.get(vehicle_registration="FLY098")
        self.assertEqual(b.fault, Claim.Fault.OUR_DRIVER)          # OI
        self.assertEqual(b.report_type, Claim.ReportType.ETARS)
        self.assertEqual(b.police_report_number, "Etars - 307250")  # original kept

    def test_dry_run_saves_nothing(self):
        import tempfile
        from django.core.management import call_command
        from .models import Claim

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        self._make_xlsx(path)
        call_command("import_accidents", path, "--dry-run")
        self.assertEqual(Claim.objects.count(), 0)
